"""
Tool: get_next_tasks  [Sonia's Agent]

After a task completion event, determines which task is now ready to start
and automatically sets it to "In progress" in state and in the Excel workbook.
Sequencing is by ascending Item ID order — no dependency column is used.
Respects the Not Ok blocker — returns empty result with blocked=True if active.
"""
import logging

import openpyxl

import state

logger = logging.getLogger(__name__)


def _auto_color():
    try:
        from tools.color_excel import color_excel
        color_excel()
    except Exception as exc:
        logger.warning("Auto color_excel after status change failed: %s", exc)


def get_next_tasks(completed_item_id: str) -> dict:
    """
    Identify the next task after a completed task and set it to "In progress".

    Automatically updates the next task's status to "In progress" in memory
    and writes it back to the Excel workbook.

    Args:
        completed_item_id: The Item ID that was just completed.

    Returns:
        A dict with next_tasks, blocked, blocking_task_id, remaining_tasks_count.
    """
    if not state.is_loaded():
        return {
            "success": False,
            "error": "No cutover plan loaded.",
            "next_tasks": [],
            "blocked": False,
            "blocking_task_id": None,
            "remaining_tasks_count": 0,
            "completed_item_id": completed_item_id,
        }

    if state.blocked_task_id:
        logger.warning("M3.missed: Blocker active — task_id=%s", state.blocked_task_id)
        return {
            "success": True,
            "next_tasks": [],
            "blocked": True,
            "blocking_task_id": state.blocked_task_id,
            "remaining_tasks_count": _count_remaining(),
            "completed_item_id": completed_item_id,
        }

    tasks = state.cutover_plan  # sorted by Item ID ascending
    remaining = [t for t in tasks if t["status"] not in ("Completed",)]

    if not remaining:
        logger.info("M3.missed: All tasks completed")
        return {
            "success": True,
            "next_tasks": [],
            "blocked": False,
            "blocking_task_id": None,
            "remaining_tasks_count": 0,
            "completed_item_id": completed_item_id,
        }

    next_task = remaining[0]
    next_id = next_task["item_id"]

    # Automatically set the next task to "In progress"
    previous_status = next_task["status"]
    next_task["status"] = "In progress"
    _write_status_to_excel(next_id, "In progress")
    _auto_color()

    logger.info(
        "M3.achieved: Next task identified and set to In progress — task_id=%s (was: %s)",
        next_id, previous_status,
    )
    return {
        "success": True,
        "next_tasks": [next_task],
        "blocked": False,
        "blocking_task_id": None,
        "remaining_tasks_count": len(remaining),
        "completed_item_id": completed_item_id,
        "auto_started": next_id,
    }


def _write_status_to_excel(item_id: str, status: str) -> None:
    path = state.workbook_path
    if not path:
        logger.warning("workbook_path not set — skipping Excel write-back for auto-start")
        return
    try:
        wb = openpyxl.load_workbook(path)
        ws = wb["SUM DMO"]
        headers = [cell.value for cell in ws[1]]
        item_id_col = headers.index("Item ID") + 1
        status_col = headers.index("Status") + 1
        for row in ws.iter_rows(min_row=2):
            if str(row[item_id_col - 1].value).strip() == item_id:
                row[status_col - 1].value = status
                break
        wb.save(path)
        logger.info("Auto-start write-back: item_id=%s status=%s", item_id, status)
    except Exception as exc:
        logger.error("Auto-start write-back failed for item_id=%s: %s", item_id, exc)


def _count_remaining() -> int:
    if not state.cutover_plan:
        return 0
    return sum(1 for t in state.cutover_plan if t["status"] not in ("Completed",))
