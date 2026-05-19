"""
Tool: mark_task_complete  [Sonia's Agent]

Updates the in-memory task status for a given Item ID and writes the change
back to the Excel workbook so it persists across sessions.
"""
import logging
from datetime import datetime, timezone

import openpyxl

import state


def _auto_color():
    try:
        from tools.color_excel import color_excel
        color_excel()
    except Exception as exc:
        logger.warning("Auto color_excel after status change failed: %s", exc)

logger = logging.getLogger(__name__)

VALID_STATUSES = {"Open", "In progress", "Completed", "Not Ok"}


def _write_status_to_excel(item_id: str, status: str) -> None:
    path = state.workbook_path
    if not path:
        logger.warning("workbook_path not set — skipping Excel write-back")
        return
    try:
        wb = openpyxl.load_workbook(path)
        ws = wb["SUM DMO"]
        headers = [cell.value for cell in ws[1]]
        item_id_col = headers.index("Item ID") + 1
        status_col = headers.index("Status") + 1
        for row in ws.iter_rows(min_row=2):
            if str(row[item_id_col - 1].value).strip() == item_id:
                row[status_col - 1].value = status
                break
        wb.save(path)
        logger.info("Excel write-back: item_id=%s status=%s", item_id, status)
    except Exception as exc:
        logger.error("Excel write-back failed for item_id=%s: %s", item_id, exc)


def mark_task_complete(item_id: str, status: str = "Completed") -> dict:
    """
    Update the status of a cutover task identified by item_id.

    Args:
        item_id: The Item ID from the SUM DMO sheet.
        status:  New status — must be one of: Open, In progress, Completed, Not Ok.

    Returns:
        A dict with item_id, new_status, trigger, next_action, and optional error.
    """
    if not state.is_loaded():
        return {
            "success": False,
            "error": "No cutover plan loaded. Please upload the workbook first.",
            "item_id": item_id,
        }

    if status not in VALID_STATUSES:
        msg = f"Invalid status '{status}'. Must be one of: {sorted(VALID_STATUSES)}"
        return {"success": False, "error": msg, "item_id": item_id}

    task = state.task_index.get(item_id)
    if task is None:
        msg = f"Item ID '{item_id}' not found in the loaded cutover plan."
        return {"success": False, "error": msg, "item_id": item_id}

    previous_status = task["status"]
    task["status"] = status
    task["completed_at"] = datetime.now(timezone.utc).isoformat() if status == "Completed" else ""
    _write_status_to_excel(item_id, status)
    _auto_color()

    if status == "Completed":
        if state.blocked_task_id == item_id:
            state.blocked_task_id = None
        logger.info("M2.achieved: Task completion detected — task_id=%s", item_id)
        return {
            "success": True,
            "item_id": item_id,
            "previous_status": previous_status,
            "new_status": status,
            "trigger": "conversation",
            "next_action": "get_next_tasks",
        }

    if status == "Not Ok":
        state.blocked_task_id = item_id
        logger.warning("M2.achieved: Task marked Not Ok — task_id=%s, blocker set", item_id)
        return {
            "success": True,
            "item_id": item_id,
            "previous_status": previous_status,
            "new_status": status,
            "trigger": "conversation",
            "next_action": "alert_pm",
            "blocker_message": (
                f"⚠️ Task {item_id} has been flagged as Not Ok. "
                "No further notifications will be sent until this is resolved."
            ),
        }

    logger.info("M2.achieved: Task status updated — task_id=%s, new_status=%s", item_id, status)
    return {
        "success": True,
        "item_id": item_id,
        "previous_status": previous_status,
        "new_status": status,
        "trigger": "conversation",
        "next_action": "none",
    }
